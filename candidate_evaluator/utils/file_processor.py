"""File processing utilities for candidate materials"""

import os
from pathlib import Path
from typing import List, Dict, Optional
import logging

# PDF processing
try:
    from PyPDF2 import PdfReader
except ImportError:
    PdfReader = None

# DOCX processing
try:
    from docx import Document
except ImportError:
    Document = None

# Markdown processing
try:
    import markdown
except ImportError:
    markdown = None


logger = logging.getLogger(__name__)


class FileProcessor:
    """Process various file formats and extract text content"""

    SUPPORTED_FORMATS = {
        '.pdf': 'process_pdf',
        '.docx': 'process_docx',
        '.txt': 'process_txt',
        '.md': 'process_markdown',
    }

    def __init__(self, max_file_size_mb: int = 10):
        """
        Initialize file processor.

        Args:
            max_file_size_mb: Maximum file size in megabytes
        """
        self.max_file_size_bytes = max_file_size_mb * 1024 * 1024

    def process_file(self, file_path: str) -> Dict[str, str]:
        """
        Process a file and extract its text content.

        Args:
            file_path: Path to the file

        Returns:
            Dictionary with 'content' and 'metadata'

        Raises:
            FileNotFoundError: If file doesn't exist
            ValueError: If file format not supported or file too large
        """
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        # Check file size
        file_size = path.stat().st_size
        if file_size > self.max_file_size_bytes:
            raise ValueError(
                f"File too large: {file_size / 1024 / 1024:.2f}MB "
                f"(max: {self.max_file_size_bytes / 1024 / 1024:.2f}MB)"
            )

        # Get file extension
        ext = path.suffix.lower()
        if ext not in self.SUPPORTED_FORMATS:
            raise ValueError(
                f"Unsupported file format: {ext}. "
                f"Supported formats: {', '.join(self.SUPPORTED_FORMATS.keys())}"
            )

        # Process based on file type
        processor_method = getattr(self, self.SUPPORTED_FORMATS[ext])
        content = processor_method(path)

        return {
            'content': content,
            'metadata': {
                'filename': path.name,
                'file_path': str(path),
                'file_type': ext,
                'file_size': file_size,
            }
        }

    def process_pdf(self, file_path: Path) -> str:
        """Extract text from PDF file"""
        if PdfReader is None:
            raise ImportError("PyPDF2 not installed. Install with: pip install PyPDF2")

        try:
            reader = PdfReader(str(file_path))
            text_parts = []

            for page_num, page in enumerate(reader.pages, 1):
                text = page.extract_text()
                if text.strip():
                    text_parts.append(f"--- Page {page_num} ---\n{text}")

            content = "\n\n".join(text_parts)
            logger.info(f"Extracted {len(content)} characters from PDF: {file_path.name}")
            return content

        except Exception as e:
            logger.error(f"Error processing PDF {file_path}: {e}")
            raise ValueError(f"Failed to process PDF: {e}")

    def process_docx(self, file_path: Path) -> str:
        """Extract text from DOCX file"""
        if Document is None:
            raise ImportError("python-docx not installed. Install with: pip install python-docx")

        try:
            doc = Document(str(file_path))
            text_parts = []

            # Extract paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)

            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        text_parts.append(" | ".join(row_text))

            content = "\n\n".join(text_parts)
            logger.info(f"Extracted {len(content)} characters from DOCX: {file_path.name}")
            return content

        except Exception as e:
            logger.error(f"Error processing DOCX {file_path}: {e}")
            raise ValueError(f"Failed to process DOCX: {e}")

    def process_txt(self, file_path: Path) -> str:
        """Extract text from TXT file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()

            logger.info(f"Extracted {len(content)} characters from TXT: {file_path.name}")
            return content

        except UnicodeDecodeError:
            # Try with different encoding
            try:
                with open(file_path, 'r', encoding='latin-1') as f:
                    content = f.read()
                logger.info(f"Extracted {len(content)} characters from TXT (latin-1): {file_path.name}")
                return content
            except Exception as e:
                logger.error(f"Error processing TXT {file_path}: {e}")
                raise ValueError(f"Failed to process TXT: {e}")

        except Exception as e:
            logger.error(f"Error processing TXT {file_path}: {e}")
            raise ValueError(f"Failed to process TXT: {e}")

    def process_markdown(self, file_path: Path) -> str:
        """Extract text from Markdown file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()

            # Keep the markdown as-is, which is readable
            # Optionally convert to plain text
            logger.info(f"Extracted {len(content)} characters from MD: {file_path.name}")
            return content

        except Exception as e:
            logger.error(f"Error processing MD {file_path}: {e}")
            raise ValueError(f"Failed to process MD: {e}")

    def process_multiple_files(self, file_paths: List[str]) -> List[Dict[str, any]]:
        """
        Process multiple files.

        Args:
            file_paths: List of file paths

        Returns:
            List of dictionaries with content and metadata

        Raises:
            ValueError: If any file fails to process
        """
        results = []
        errors = []

        for file_path in file_paths:
            try:
                result = self.process_file(file_path)
                results.append(result)
            except Exception as e:
                error_msg = f"Failed to process {file_path}: {e}"
                logger.error(error_msg)
                errors.append(error_msg)

        if errors:
            raise ValueError(f"Errors processing files:\n" + "\n".join(errors))

        return results

    def combine_materials(self, processed_files: List[Dict[str, any]]) -> str:
        """
        Combine multiple processed files into a single text.

        Args:
            processed_files: List of processed file dictionaries

        Returns:
            Combined text with clear section markers
        """
        sections = []

        for file_data in processed_files:
            filename = file_data['metadata']['filename']
            content = file_data['content']

            section = f"""
{'='*80}
DOCUMENT: {filename}
{'='*80}

{content}
"""
            sections.append(section)

        return "\n\n".join(sections)
